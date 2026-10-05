# -*- coding: utf-8 -*-
"""Формирование отчёта по лабораторной работе № 3.

Оформление полностью повторяет отчёты № 1 и № 2:
A4, поля 30/15/20/20 мм, Times New Roman 14 пт, интервал 1,5,
абзацный отступ 1,25 см, выравнивание по ширине, нумерация страниц
снизу по центру (первая страница без номера).
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.shared import Cm, Pt

import format_gost as fg
from format_gost import (BODY, CENTER, FIG_CAP, HEAD, MONOBLOCK, MONO, TNR,
                         fill, fmt_para, new_para, setup_section, setup_styles,
                         table_format)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Лабораторная_3_Настройка_фильтров_импорта.docx"

PLACEHOLDER = dict(align=AL.CENTER, first=Cm(0), ls=1.5, before=6, after=6)
LISTING = dict(align=AL.LEFT, first=Cm(0), ls=1.0, before=6, after=6)


def h1(doc, text):
    return new_para(doc, text, HEAD, bold=True)


def listing(doc, text, mono=True):
    p = doc.add_paragraph()
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        p.add_run(line)
    return fmt_para(p, LISTING, name=MONO if mono else TNR, size=12)


def add_table(doc, data, caption, widths=None):
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    t.style = "Table Grid"
    for ri, row in enumerate(data):
        for ci, value in enumerate(row):
            t.cell(ri, ci).text = value
    table_format(doc, len(doc.tables) - 1, caption)
    return t


def placeholder(doc, what):
    return new_para(doc, f"Место для скриншота:\n\n[ ВСТАВИТЬ СКРИНШОТ: {what} ]",
                    PLACEHOLDER, bold=True)


def figure(doc, caption):
    return new_para(doc, caption, FIG_CAP)


def build():
    doc = Document()
    setup_styles(doc)
    setup_section(doc)

    # ------------------------------------------------ титульная часть
    for text in ("ЛАБОРАТОРНАЯ РАБОТА № 3",
                 "Настройка работы системы контроля версий",
                 "(типов импортируемых файлов, путей, фильтров и других "
                 "параметров импорта в репозиторий)"):
        new_para(doc, text, dict(CENTER, after=6), bold=True)
    new_para(doc, "", CENTER)
    new_para(doc, "Тема проекта: «Генератор случайных цитат»",
             dict(CENTER, after=6), bold=True)
    new_para(doc, "", CENTER)
    new_para(doc, "Студент: Сагадиев Амир\nГруппа: ИСП-23-19", CENTER)
    new_para(doc, "Преподаватель: Кондров В. А.", CENTER)
    new_para(doc, "", CENTER)

    # ------------------------------------------------ 1. теория
    h1(doc, "1. Краткие теоретические сведения")
    new_para(doc, "Система контроля версий хранит историю изменений проекта "
                  "и позволяет вернуться к любой из предыдущих версий. При "
                  "этом в репозиторий должны попадать только файлы, "
                  "необходимые для работы над проектом: исходный код, данные, "
                  "документация и настройки. Служебные, автоматически "
                  "создаваемые и конфиденциальные файлы импортироваться не "
                  "должны, поэтому их импорт ограничивается специальными "
                  "настройками Git.", BODY)
    new_para(doc, "Файл .gitignore содержит список шаблонов имён файлов и "
                  "каталогов, которые Git не отслеживает. Шаблоны вида "
                  "__pycache__/ или *.pyc описывают целые группы файлов, а "
                  "знак «!» отменяет ранее заданное исключение. Правила "
                  ".gitignore применяются к файлам, которые ещё не добавлены "
                  "в индекс; если файл уже отслеживается, его предварительно "
                  "убирают из индекса командой git rm --cached.", BODY)
    new_para(doc, "Файл .gitattributes задаёт атрибуты путей и определяет, "
                  "как Git обрабатывает содержимое файлов. Основные "
                  "возможности: приведение переводов строк к единому формату "
                  "(* text=auto eol=lf), пометка бинарных файлов, чтобы Git не "
                  "пытался сравнивать в них строки (*.png binary), а также "
                  "привязка больших файлов к системе Git LFS.", BODY)
    new_para(doc, "Git LFS (Large File Storage) — расширение Git для хранения "
                  "больших файлов: изображений, архивов, дизайн-макетов. Сами "
                  "файлы хранятся в отдельном хранилище, а в репозиторий "
                  "попадает только небольшой текстовый указатель. Утилита "
                  "устанавливается в систему, активируется командой "
                  "git lfs install, после чего типы файлов привязываются "
                  "командой git lfs track, например git lfs track \"*.psd\".", BODY)
    new_para(doc, "Сервис gitignore.io (Toptal) — онлайн-подсказка, которая "
                  "формирует готовый шаблон .gitignore по списку используемых "
                  "технологий и сред разработки. В данной работе шаблон "
                  "сгенерирован для набора python, pycharm, windows, "
                  "visualstudiocode, visualstudio и затем дополнен разделами "
                  "конфиденциальных и временных файлов.", BODY)

    # ------------------------------------------------ 2. перечень фильтров
    h1(doc, "2. Перечень фильтров импорта")
    new_para(doc, "В соответствии с заданием настраиваются три группы правил: "
                  "фильтры импорта файлов и каталогов (.gitignore), типы "
                  "файлов и правила работы с ними (.gitattributes), а также "
                  "импорт тяжёлых ресурсов через Git LFS. Перечень "
                  "исключаемых файлов приведён в таблице 1.", BODY)
    add_table(doc, [
        ["Категория", "Примеры", "Причина исключения"],
        ["Служебные файлы среды разработки",
         ".idea/, .vscode/, *.iml, *.iws, Debug/, Release/, *.user",
         "Содержат локальные настройки конкретного компьютера и не нужны "
         "другим участникам проекта"],
        ["Зависимости и кэш",
         "__pycache__/, *.pyc, venv/, .venv/, node_modules/, .target/",
         "Создаются автоматически и могут быть восстановлены; увеличивают "
         "размер репозитория"],
        ["Конфиденциальные данные",
         ".env, .env.*, secrets.json, credentials.json, *.pem, *.key",
         "Секреты не должны попадать в историю Git и публиковаться"],
        ["Системные и временные файлы",
         ".DS_Store, Thumbs.db, *.log, *.tmp, ~$*",
         "Создаются операционной системой и редакторами, к проекту не "
         "относятся"],
    ], "Таблица 1 – Файлы, исключаемые из репозитория")
    new_para(doc, "Файл config.json намеренно не исключается: в нём хранятся "
                  "только несекретные настройки (путь к файлу цитат и язык "
                  "интерфейса). Конфиденциальные параметры, согласно отчёту "
                  "№ 2, должны храниться в файле .env, который добавлен в "
                  ".gitignore.", BODY)

    # ------------------------------------------------ 3. .gitignore
    h1(doc, "3. Создание файла .gitignore")
    h1(doc, "3.1. Состав фильтров")
    new_para(doc, "Файл .gitignore создан в корне репозитория и состоит из "
                  "четырёх разделов:", BODY)
    for line in ("1. Служебные файлы среды разработки (PyCharm, VS Code, "
                 "Visual Studio, прочие редакторы).",
                 "2. Зависимости и кэш: кэш байт-кода Python, виртуальные "
                 "окружения, кэш инструментов, каталоги сборки.",
                 "3. Конфиденциальные данные: файлы .env, секреты, ключи и "
                 "сертификаты.",
                 "4. Системные и временные файлы операционной системы и "
                 "редакторов."):
        new_para(doc, line, BODY)
    new_para(doc, "Правила действуют на все вложенные каталоги, включая "
                  "каталог проекта lab1_quote_generator.", BODY)

    h1(doc, "3.2. Листинг файла .gitignore")
    listing(doc, (ROOT / ".gitignore").read_text(encoding="utf-8"))

    h1(doc, "3.3. Проверка фильтров")
    new_para(doc, "Проверка выполняется командой git check-ignore -v: она "
                  "показывает, какое правило и в какой строке .gitignore "
                  "исключает файл.", BODY)
    listing(doc,
            "$ git check-ignore -v "
            "lab1_quote_generator/src/modules/__pycache__/ui.cpython-314.pyc\n"
            ".gitignore:41:__pycache__/\t"
            "lab1_quote_generator/src/modules/__pycache__/ui.cpython-314.pyc\n"
            "\n"
            "$ git check-ignore -v .idea/workspace.xml\n"
            ".gitignore:14:.idea/\t.idea/workspace.xml")
    new_para(doc, "Ранее добавленный в репозиторий кэш байт-кода убран из "
                  "индекса командой git rm -r --cached "
                  "lab1_quote_generator/src/modules/__pycache__: сами файлы "
                  "остались на диске, но Git их больше не отслеживает.", BODY)

    # ------------------------------------------------ 4. .gitattributes
    h1(doc, "4. Настройка типов файлов и путей (.gitattributes)")
    h1(doc, "4.1. Назначение атрибутов")
    new_para(doc, "Файл .gitattributes создан в корне репозитория и "
                  "настраивает три группы правил: единый формат переводов "
                  "строк для текстовых файлов, пометку бинарных файлов и "
                  "привязку тяжёлых ресурсов к Git LFS (таблица 2).", BODY)
    add_table(doc, [
        ["Шаблон пути", "Атрибут", "Назначение"],
        ["*", "text=auto eol=lf",
         "Все текстовые файлы приводятся к единому формату переноса строки LF"],
        ["*.py, *.json, *.md, *.txt, *.yml", "text eol=lf",
         "Явное правило для исходного кода и файлов данных"],
        ["*.png, *.jpg, *.jpeg, *.gif, *.ico, *.bmp", "binary",
         "Файлы помечены как бинарные: Git не сравнивает в них строки"],
        ["*.pdf, *.docx, *.xlsx, *.pptx, *.exe, *.dll", "binary",
         "Документы и исполняемые файлы также помечены как бинарные"],
        ["*.psd, *.zip", "filter=lfs diff=lfs merge=lfs -text",
         "Файлы импортируются в репозиторий через Git LFS"],
    ], "Таблица 2 – Атрибуты путей, заданные в .gitattributes")

    h1(doc, "4.2. Листинг файла .gitattributes")
    listing(doc, (ROOT / ".gitattributes").read_text(encoding="utf-8"))

    h1(doc, "4.3. Проверка атрибутов")
    listing(doc,
            "$ git check-attr -a -- assets/example_design.psd\n"
            "assets/example_design.psd: diff: lfs\n"
            "assets/example_design.psd: merge: lfs\n"
            "assets/example_design.psd: text: unset\n"
            "assets/example_design.psd: filter: lfs\n"
            "\n"
            "$ git check-attr -a -- lab1_quote_generator/data/quotes.json\n"
            "lab1_quote_generator/data/quotes.json: text: set\n"
            "lab1_quote_generator/data/quotes.json: eol: lf\n"
            "\n"
            "$ git check-attr -a -- docs/diagram.png\n"
            "docs/diagram.png: binary: set\n"
            "docs/diagram.png: diff: unset\n"
            "docs/diagram.png: merge: unset\n"
            "\n"
            "$ git ls-files --eol lab1_quote_generator/src/main.py\n"
            "i/lf    w/lf    attr/text eol=lf\t"
            "lab1_quote_generator/src/main.py")
    new_para(doc, "Атрибут text eol=lf подтверждён для файла src/main.py: "
                  "файл хранится в репозитории и выгружается с переводами "
                  "строк LF независимо от операционной системы.", BODY)

    # ------------------------------------------------ 5. Git LFS
    h1(doc, "5. Настройка импорта тяжёлых ресурсов (Git LFS)")
    h1(doc, "5.1. Установка и активация утилиты")
    new_para(doc, "Утилита Git LFS скачивается с официального сайта "
                  "git-lfs.com (установщик для Windows) либо устанавливается "
                  "пакетным менеджером winget install --id GitHub.GitLFS. "
                  "После установки утилита активируется командой:", BODY)
    listing(doc, "$ git lfs install\nUpdated Git hooks.\nGit LFS initialized.")
    new_para(doc, "Команда прописывает в конфигурацию Git фильтры LFS и "
                  "устанавливает служебные обработчики событий репозитория.", BODY)

    h1(doc, "5.2. Привязка типов файлов к LFS")
    new_para(doc, "Тяжёлые типы файлов привязываются к LFS командами:", BODY)
    listing(doc, "$ git lfs track \"*.psd\"\nTracking \"*.psd\"\n"
                 "$ git lfs track \"*.zip\"\nTracking \"*.zip\"")
    new_para(doc, "Команды автоматически добавляют в конец файла "
                  ".gitattributes соответствующие строки:", BODY)
    listing(doc, "*.psd filter=lfs diff=lfs merge=lfs -text\n"
                 "*.zip filter=lfs diff=lfs merge=lfs -text")
    new_para(doc, "Таким образом, при добавлении в репозиторий файлов с "
                  "расширениями .psd и .zip они будут импортироваться через "
                  "Git LFS: в основной истории сохранится только текстовый "
                  "указатель, а сам файл попадёт в отдельное хранилище.", BODY)

    h1(doc, "5.3. Проверка импорта")
    new_para(doc, "Список файлов, импортированных через LFS, выводится "
                  "командой git lfs ls-files. Пока в проекте нет тяжёлых "
                  "файлов, поэтому список пуст. При добавлении файла с "
                  "расширением .psd или .zip он появится в этом списке, а его "
                  "строки в .gitattributes будут иметь вид filter=lfs "
                  "diff=lfs merge=lfs -text.", BODY)

    # ------------------------------------------------ 6. проверка
    h1(doc, "6. Результаты проверки настроек")
    new_para(doc, "Сводные результаты проверки настройки фильтров и атрибутов "
                  "приведены в таблице 3.", BODY)
    add_table(doc, [
        ["Что проверялось", "Команда", "Результат"],
        ["Кэш байт-кода игнорируется",
         "git check-ignore -v .../__pycache__/ui.cpython-314.pyc",
         "Файл исключён правилом .gitignore:41 (__pycache__/)"],
        ["Служебные файлы IDE игнорируются",
         "git check-ignore -v .idea/workspace.xml",
         "Файл исключён правилом .gitignore:14 (.idea/)"],
        ["Исходный код не исключается",
         "git check-ignore -v lab1_quote_generator/src/main.py",
         "Правило не найдено: файл отслеживается репозиторием"],
        ["Переводы строк текстовых файлов",
         "git ls-files --eol lab1_quote_generator/src/main.py",
         "i/lf w/lf attr/text eol=lf"],
        ["Бинарные файлы помечены",
         "git check-attr -a -- docs/diagram.png",
         "binary: set, diff: unset, merge: unset"],
        ["Тяжёлые ресурсы привязаны к LFS",
         "git check-attr -a -- assets/example_design.psd",
         "filter: lfs, diff: lfs, merge: lfs, text: unset"],
        ["Кэш удалён из индекса",
         "git rm -r --cached .../__pycache__",
         "Файлы убраны из индекса, на диске сохранены"],
    ], "Таблица 3 – Проверка настройки фильтров и атрибутов")

    # ------------------------------------------------ 7. результат
    h1(doc, "7. Результат работы")
    new_para(doc, "В ходе лабораторной работы настроены параметры импорта "
                  "файлов в репозиторий проекта «Генератор случайных цитат». "
                  "Создан файл .gitignore, исключающий служебные файлы сред "
                  "разработки, зависимости и кэш, конфиденциальные данные, а "
                  "также системные и временные файлы. Создан файл "
                  ".gitattributes, который приводит переводы строк текстовых "
                  "файлов к единому формату LF, помечает бинарные файлы и "
                  "привязывает тяжёлые ресурсы к Git LFS. Утилита Git LFS "
                  "активирована, типы файлов .psd и .zip привязаны к ней. "
                  "Ранее добавленный в репозиторий кэш байт-кода убран из "
                  "индекса. Настроенные правила проверены командами "
                  "git check-ignore, git check-attr и git ls-files --eol.", BODY)
    new_para(doc, "Настроенная конфигурация обеспечивает чистоту истории "
                  "репозитория и может использоваться в следующих "
                  "лабораторных работах.", BODY)

    # ------------------------------------------------ 8. рисунки
    h1(doc, "8. Контрольные материалы")
    figures = [
        ("На рисунке 1 представлено дерево репозитория с созданными файлами "
         ".gitignore и .gitattributes в корне.",
         "окно Project в PyCharm: корень репозитория, файлы .gitignore и "
         ".gitattributes",
         "Рисунок 1 – дерево репозитория с файлами .gitignore и "
         ".gitattributes"),
        ("На рисунке 2 представлено содержимое файла .gitignore.",
         "открытый в редакторе PyCharm файл .gitignore (видны разделы "
         "фильтров)",
         "Рисунок 2 – содержимое файла .gitignore"),
        ("На рисунке 3 представлено содержимое файла .gitattributes.",
         "открытый в редакторе PyCharm файл .gitattributes, включая строки "
         "Git LFS (*.psd, *.zip)",
         "Рисунок 3 – содержимое файла .gitattributes"),
        ("На рисунке 4 представлена проверка фильтров командой "
         "git check-ignore.",
         "терминал PyCharm: команды git check-ignore -v для файла кэша и "
         "служебного файла IDE",
         "Рисунок 4 – проверка фильтров командой git check-ignore"),
        ("На рисунке 5 представлена настройка Git LFS и проверка привязки "
         "типов файлов.",
         "терминал PyCharm: команды git lfs install, git lfs track \"*.psd\", "
         "git lfs track \"*.zip\" и git lfs ls-files",
         "Рисунок 5 – настройка Git LFS и привязка типов файлов"),
    ]
    for text, what, caption in figures:
        new_para(doc, text, BODY)
        placeholder(doc, what)
        figure(doc, caption)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print("готово:", path)

# -*- coding: utf-8 -*-
"""Формирование отчёта по лабораторной работе № 4.

Оформление полностью повторяет отчёты № 1–3:
A4, поля 30/15/20/20 мм, Times New Roman 14 пт, интервал 1,5,
абзацный отступ 1,25 см, выравнивание по ширине, нумерация страниц
снизу по центру (первая страница без номера).
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.shared import Cm

from format_gost import (BODY, CENTER, FIG_CAP, HEAD, MONO,
                         fmt_para, new_para, setup_section, setup_styles,
                         table_format)

ROOT = Path(__file__).resolve().parent.parent
PROJECT = ROOT / "lab1_quote_generator"
OUT = ROOT / "Лабораторная_4_Разработка_и_интеграция_модулей.docx"

PLACEHOLDER = dict(align=AL.CENTER, first=Cm(0), ls=1.5, before=6, after=6)
LISTING = dict(align=AL.LEFT, first=Cm(0), ls=1.0, before=6, after=6)


def read(path):
    return (PROJECT / path).read_text(encoding="utf-8")


def extract(path, names):
    """Вырезать из файла функции с указанными именами (для фрагментов)."""
    text = read(path)
    blocks = re.split(r"\n(?=def )", text)
    picked = []
    for block in blocks:
        for name in names:
            if block.startswith(f"def {name}("):
                picked.append(block.rstrip("\n"))
    return "\n\n\n".join(picked)


def h1(doc, text):
    return new_para(doc, text, HEAD, bold=True)


def listing(doc, text, size=12):
    p = doc.add_paragraph()
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        p.add_run(line)
    return fmt_para(p, LISTING, name=MONO, size=size)


def add_table(doc, data, caption):
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    t.style = "Table Grid"
    for ri, row in enumerate(data):
        for ci, value in enumerate(row):
            t.cell(ri, ci).text = value
    table_format(doc, len(doc.tables) - 1, caption)
    return t


def placeholder(doc, what):
    return new_para(doc, f"Место для скриншота:\n\n[ ВСТАВИТЬ СКРИНШОТ: {what} ]",
                    PLACEHOLDER, bold=True)


def figure(doc, caption):
    return new_para(doc, caption, FIG_CAP)


def build():
    doc = Document()
    setup_styles(doc)
    setup_section(doc)

    # ------------------------------------------------ титульная часть
    new_para(doc, "ЛАБОРАТОРНАЯ РАБОТА № 4", dict(CENTER, after=6), bold=True)
    new_para(doc, "Разработка и интеграция модулей проекта",
             dict(CENTER, after=6), bold=True)
    new_para(doc, "", CENTER)
    new_para(doc, "Тема проекта: «Генератор случайных цитат»",
             dict(CENTER, after=6), bold=True)
    new_para(doc, "", CENTER)
    new_para(doc, "Студент: Сагадиев Амир\nГруппа: ИСП-23-19", CENTER)
    new_para(doc, "Преподаватель: Кондров В. А.", CENTER)
    new_para(doc, "", CENTER)

    # ------------------------------------------------ 1. теория
    h1(doc, "1. Краткие теоретические сведения")
    new_para(doc, "Модуль — это отдельная часть программы, которая отвечает "
                  "за одну задачу и взаимодействует с остальными частями "
                  "через понятный интерфейс: определённые входные данные и "
                  "определённый результат. Модуль считается независимым, "
                  "если он не знает о внутреннем устройстве других модулей: "
                  "например, модуль выбора цитаты работает со списком, "
                  "полученным извне, и не читает файлы самостоятельно.", BODY)
    new_para(doc, "Интеграция — это объединение готовых модулей в единое "
                  "приложение. Для этого пишется связующий код, который "
                  "передаёт результат работы одного модуля на вход другому. "
                  "Чаще всего таким связующим звеном является файл точки "
                  "входа (main.py).", BODY)
    new_para(doc, "Разработка каждого модуля ведётся в отдельной ветке Git "
                  "(feature-ветка), что позволяет не мешать основную ветку "
                  "незаконченным кодом. После разработки модулей создаётся "
                  "интеграционная ветка, в которой выполняется сборка и "
                  "проверка совместной работы, а затем ветка вливается в "
                  "основную.", BODY)
    new_para(doc, "Качество модулей подтверждается тестами. Модульные "
                  "(юнит) тесты проверяют каждый модуль отдельно, локально, "
                  "без запуска всего приложения. Интеграционное тестирование "
                  "проверяет сквозной сценарий: от входных данных до "
                  "получения итогового результата.", BODY)

    # ------------------------------------------------ 2. архитектура
    h1(doc, "2. Проектирование архитектуры")
    new_para(doc, "Минимальный функционал (MVP) приложения «Генератор "
                  "случайных цитат»: загрузить набор цитат из файла и "
                  "вывести одну случайную цитату. Для его работы выделены "
                  "два независимых модуля (таблица 1).", BODY)
    add_table(doc, [
        ["Модуль", "Файл и класс", "Вход", "Выход", "Зависимости"],
        ["№ 1. Парсер данных",
         "src/modules/quote_storage.py, класс QuoteStorage",
         "Путь к файлу data/quotes.json",
         "Список цитат: [{\"text\": ..., \"category\": ...}, ...]",
         "Нет (работает только с файлом)"],
        ["№ 2. Генератор цитат",
         "src/modules/quote_generator.py, класс QuoteGenerator",
         "Список цитат от Модуля № 1 и необязательная категория",
         "Одна случайная цитата: {\"text\": ..., \"category\": ...}",
         "Нет (работает только со списком)"],
    ], "Таблица 1 – Модули проекта и их интерфейсы")
    new_para(doc, "Схема передачи данных: Модуль № 1 читает файл и отдаёт "
                  "список цитат; связующий код в main.py передаёт этот "
                  "список на вход Модулю № 2; Модуль № 2 выбирает одну "
                  "случайную цитату, которую выводит интерфейс. Схема "
                  "зафиксирована в README проекта (раздел «Архитектура: как "
                  "модули передают данные»).", BODY)
    listing(doc,
            "                 data/quotes.json\n"
            "                        │\n"
            "                        ▼\n"
            "        ┌───────────────────────────────┐\n"
            "        │  Модуль № 1 «Парсер данных»   │\n"
            "        │  QuoteStorage.load()          │\n"
            "        └───────────────┬───────────────┘\n"
            "                        │  список цитат\n"
            "                        ▼\n"
            "        ┌───────────────────────────────┐\n"
            "        │  Модуль № 2 «Генератор»       │\n"
            "        │  QuoteGenerator               │\n"
            "        │  .random_quote(category)      │\n"
            "        └───────────────┬───────────────┘\n"
            "                        │  одна цитата\n"
            "                        ▼\n"
            "        ┌───────────────────────────────┐\n"
            "        │  main.py + ui.py              │\n"
            "        │  вывод пользователю           │\n"
            "        └───────────────────────────────┘", size=11)

    # ------------------------------------------------ 3. модуль 1
    h1(doc, "3. Разработка Модуля № 1 (парсер данных)")
    new_para(doc, "Под модуль создана отдельная ветка Git:", BODY)
    listing(doc, "$ git checkout -b feature/data-parser")
    new_para(doc, "Модуль оформлен в виде отдельного класса QuoteStorage. "
                  "Он читает JSON-файл, проверяет структуру и возвращает "
                  "список цитат, а при проблемах поднимает понятную ошибку "
                  "QuoteStorageError: файл не найден, повреждён JSON, "
                  "данные не являются списком. Записи без текста "
                  "пропускаются. Модуль не зависит от других модулей "
                  "проекта.", BODY)

    h1(doc, "3.1. Листинг модуля")
    listing(doc, read("src/modules/quote_storage.py"))

    h1(doc, "3.2. Модульные тесты")
    new_para(doc, "Модуль протестирован отдельно, без запуска всего "
                  "приложения. Фрагмент листинга тестов:", BODY)
    listing(doc, extract("tests/test_quote_storage.py", [
        "test_load_returns_all_quotes",
        "test_records_without_text_are_skipped",
        "test_missing_file_raises_error",
    ]))
    new_para(doc, "Запуск тестов и результат:", BODY)
    listing(doc,
            "$ python -m pytest tests/test_quote_storage.py -v\n"
            "...\n"
            "tests/test_quote_storage.py::test_load_returns_all_quotes PASSED\n"
            "tests/test_quote_storage.py::test_records_without_text_are_skipped PASSED\n"
            "tests/test_quote_storage.py::test_missing_file_raises_error PASSED\n"
            "...\n"
            "============================== 8 passed in 0.02s ==============================",
            size=11)
    new_para(doc, "Изменения зафиксированы коммитом:", BODY)
    listing(doc, "$ git add .\n"
                 "$ git commit -m \"feat(storage): модуль парсинга данных с валидацией\"\n"
                 "$ git push -u origin feature/data-parser")

    # ------------------------------------------------ 4. модуль 2
    h1(doc, "4. Разработка Модуля № 2 (генератор цитат)")
    new_para(doc, "Под второй модуль создана отдельная ветка Git:", BODY)
    listing(doc, "$ git checkout main\n"
                 "$ git checkout -b feature/quote-generator")
    new_para(doc, "Модуль оформлен в виде отдельного класса QuoteGenerator. "
                  "На вход он получает готовый список цитат, а возвращает "
                  "одну случайную цитату. Дополнительно поддерживается "
                  "выбор по категории и получение списка категорий. Модуль "
                  "не работает с файлами и не зависит от Модуля № 1: "
                  "список передаётся ему снаружи, поэтому модуль можно "
                  "тестировать отдельно.", BODY)

    h1(doc, "4.1. Листинг модуля")
    listing(doc, read("src/modules/quote_generator.py"))

    h1(doc, "4.2. Модульные тесты")
    new_para(doc, "Вместо прежней заглушки написаны полноценные тесты. "
                  "Фрагмент листинга:", BODY)
    listing(doc, extract("tests/test_quote_generator.py", [
        "test_random_quote_returns_item_from_list",
        "test_random_quote_filters_by_category",
        "test_unknown_category_raises_error",
    ]))
    new_para(doc, "Запуск тестов и результат:", BODY)
    listing(doc,
            "$ python -m pytest tests/test_quote_generator.py -v\n"
            "...\n"
            "tests/test_quote_generator.py::test_random_quote_returns_item_from_list PASSED\n"
            "tests/test_quote_generator.py::test_random_quote_filters_by_category PASSED\n"
            "tests/test_quote_generator.py::test_unknown_category_raises_error PASSED\n"
            "...\n"
            "============================== 8 passed in 0.01s ==============================",
            size=11)
    new_para(doc, "Изменения зафиксированы коммитом:", BODY)
    listing(doc, "$ git add .\n"
                 "$ git commit -m \"feat(generator): модуль генерации цитат с фильтром по категориям\"\n"
                 "$ git push -u origin feature/quote-generator")

    # ------------------------------------------------ 5. интеграция
    h1(doc, "5. Интеграция и сборка")
    new_para(doc, "Для сборки создана интеграционная ветка, в неё влиты "
                  "обе feature-ветки:", BODY)
    listing(doc, "$ git checkout main\n"
                 "$ git checkout -b feature/integration\n"
                 "$ git merge feature/data-parser\n"
                 "$ git merge feature/quote-generator")

    h1(doc, "5.1. Связующий код")
    new_para(doc, "Связующий код написан в файле main.py: он читает "
                  "настройки, получает список цитат из Модуля № 1 и "
                  "автоматически передаёт его на вход Модулю № 2.", BODY)
    listing(doc, read("src/main.py"))
    new_para(doc, "Интерфейс выводит цитату вместе с её категорией:", BODY)
    listing(doc, read("src/modules/ui.py"))

    h1(doc, "5.2. Интеграционный тест")
    new_para(doc, "Сквозной сценарий «от файла с данными до готовой цитаты» "
                  "закреплён интеграционным тестом:", BODY)
    listing(doc, read("tests/test_integration.py"))
    listing(doc,
            "$ python -m pytest tests/test_integration.py -v\n"
            "...\n"
            "============================== 2 passed in 0.01s ===============================",
            size=11)

    h1(doc, "5.3. Ручной сквозной сценарий")
    new_para(doc, "Выполнен ручной запуск приложения от входных данных до "
                  "готового результата:", BODY)
    listing(doc,
            "$ python src/main.py\n"
            "Генератор случайных цитат\n"
            "Категория: Мотивация\n"
            "Начать — уже половина дела.\n"
            "\n"
            "$ python src/main.py\n"
            "Генератор случайных цитат\n"
            "Категория: Учёба\n"
            "Знания становятся силой, когда применяются на практике.",
            size=11)
    new_para(doc, "Приложение выводит разные цитаты при повторных запусках, "
                  "что подтверждает случайность выбора. Чтобы проверить "
                  "фильтр по категории, в config.json добавляется строка "
                  "\"category\": \"Мотивация\" — тогда выбор будет идти "
                  "только среди цитат этой категории.", BODY)

    h1(doc, "5.4. Финальный коммит и слияние")
    new_para(doc, "Интеграционная ветка зафиксирована и влита в основную:", BODY)
    listing(doc, "$ git add .\n"
                 "$ git commit -m \"feat: интеграция модулей в main.py\"\n"
                 "$ git push -u origin feature/integration\n"
                 "\n"
                 "$ git checkout main\n"
                 "$ git merge --no-ff feature/integration -m \"merge: интеграция модулей проекта\"\n"
                 "$ git push origin main")

    # ------------------------------------------------ 6. тестирование
    h1(doc, "6. Результаты тестирования")
    new_para(doc, "Все проверки запускаются из корня проекта командой "
                  "python -m pytest -v (пути к модулям настраивает "
                  "conftest.py). Сводные результаты приведены в таблице 2.", BODY)
    add_table(doc, [
        ["Набор тестов", "Что проверяет", "Результат"],
        ["tests/test_quote_storage.py (8 тестов, Модуль № 1)",
         "Загрузка цитат, значение категории по умолчанию, пропуск записей "
         "без текста, удаление лишних пробелов, ошибки: файл не найден, "
         "повреждённый JSON, данные не список, нет корректных цитат",
         "8 passed"],
        ["tests/test_quote_generator.py (8 тестов, Модуль № 2)",
         "Выбор цитаты из списка, случайность выбора, фильтр по категории, "
         "ошибка для несуществующей категории, список категорий, пустой "
         "список, неизменность исходных данных",
         "8 passed"],
        ["tests/test_integration.py (2 теста)",
         "Сквозной сценарий: файл → Модуль № 1 → Модуль № 2 → цитата; "
         "работа на реальном файле data/quotes.json",
         "2 passed"],
        ["Все тесты вместе",
         "Модульное и интеграционное тестирование",
         "18 passed"],
    ], "Таблица 2 – Результаты тестирования")

    # ------------------------------------------------ 7. результат
    h1(doc, "7. Результат работы")
    new_para(doc, "В ходе лабораторной работы выделены два независимых "
                  "модуля, обеспечивающих минимальный функционал проекта: "
                  "Модуль № 1 «Парсер данных» (загрузка и проверка цитат из "
                  "JSON-файла) и Модуль № 2 «Генератор цитат» (случайный "
                  "выбор цитаты, в том числе по категории). Каждый модуль "
                  "разработан в отдельной ветке Git и протестирован "
                  "отдельно: 8 и 8 модульных тестов. Схема передачи данных "
                  "зафиксирована в README проекта.", BODY)
    new_para(doc, "В интеграционной ветке написан связующий код в main.py: "
                  "результат работы Модуля № 1 автоматически передаётся на "
                  "вход Модулю № 2. Проведено интеграционное тестирование — "
                  "автоматический сквозной тест (2 проверки) и ручной "
                  "запуск приложения. Всего проходит 18 тестов. "
                  "Интеграционная ветка влита в основную ветку проекта.", BODY)

    # ------------------------------------------------ 8. рисунки
    h1(doc, "8. Контрольные материалы")
    figures = [
        ("На рисунке 1 представлена структура проекта в PyCharm с "
         "созданными модулями и тестами.",
         "окно Project в PyCharm: папка src/modules с файлами "
         "quote_storage.py и quote_generator.py, папка tests с тремя "
         "файлами тестов, файл conftest.py",
         "Рисунок 1 – структура проекта с модулями и тестами"),
        ("На рисунке 2 представлен список веток репозитория.",
         "терминал PyCharm: вывод команды git branch -a — видны ветки "
         "main, feature/data-parser, feature/quote-generator",
         "Рисунок 2 – список веток репозитория"),
        ("На рисунке 3 представлен запуск модульных тестов Модуля № 1.",
         "терминал PyCharm: команда python -m pytest "
         "tests/test_quote_storage.py -v и результат 8 passed",
         "Рисунок 3 – запуск модульных тестов Модуля № 1"),
        ("На рисунке 4 представлен запуск модульных тестов Модуля № 2.",
         "терминал PyCharm: команда python -m pytest "
         "tests/test_quote_generator.py -v и результат 8 passed",
         "Рисунок 4 – запуск модульных тестов Модуля № 2"),
        ("На рисунке 5 представлен сквозной запуск приложения.",
         "терминал PyCharm: команда python src/main.py и вывод цитаты с "
         "категорией (или окно Run в PyCharm)",
         "Рисунок 5 – сквозной запуск main.py"),
        ("На рисунке 6 представлена история коммитов после слияния "
         "интеграционной ветки.",
         "терминал PyCharm: команда git log --oneline --graph --decorate "
         "--all — видны коммиты модулей и слияние ветки feature/integration",
         "Рисунок 6 – история коммитов и слияние ветки интеграции"),
    ]
    for text, what, caption in figures:
        new_para(doc, text, BODY)
        placeholder(doc, what)
        figure(doc, caption)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print("готово:", path)
